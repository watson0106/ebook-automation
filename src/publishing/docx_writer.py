"""python-docx を使って GeneratedBook を .docx ファイルに書き出すモジュール

生成した .docx を Google ドライブにアップロードすると、
Google ドキュメントとして自動変換されて編集できます。
"""
import re
from pathlib import Path
from typing import Optional

from rich.console import Console

from config.settings import settings
from src.content import GeneratedBook

console = Console()


def _add_markdown_paragraphs(doc, text: str):
    """マークダウンのテキストを解析して docx に段落として追加する"""
    for line in text.split("\n"):
        # 見出し
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            # Heading 1〜6 に対応（docx は Heading 1〜9 をサポート）
            doc.add_heading(m.group(2).strip(), level=level)
            continue

        # 空行はスキップ
        if line.strip() == "":
            continue

        # 通常テキスト（**bold** のみ対応）
        para = doc.add_paragraph()
        parts = re.split(r"(\*\*[^*]+\*\*)", line)
        for part in parts:
            bold_m = re.match(r"\*\*([^*]+)\*\*", part)
            if bold_m:
                para.add_run(bold_m.group(1)).bold = True
            else:
                para.add_run(part)


def build_docx(book: GeneratedBook, output_dir: Optional[Path] = None) -> Path:
    """
    GeneratedBook を .docx ファイルに書き出す。

    Args:
        book: 書き出す本のデータ
        output_dir: 出力先ディレクトリ（省略時は settings.output_epubs_dir と同じディレクトリ）

    Returns:
        生成した .docx ファイルのパス
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as e:
        raise RuntimeError(
            "python-docx が見つかりません。'pip install python-docx' を実行してください。"
        ) from e

    out_dir = output_dir or settings.output_epubs_dir
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_title = re.sub(r'[\\/:*?"<>|]', "_", book.book_title)
    out_path = out_dir / f"{safe_title}.docx"

    doc = Document()

    # ─── タイトル ───
    title_para = doc.add_heading(book.book_title, level=0)  # level=0 → "Title" スタイル
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if book.subtitle:
        sub_para = doc.add_paragraph(book.subtitle)
        sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in sub_para.runs:
            run.font.size = Pt(14)
            run.font.italic = True

    # ─── 書誌情報 ───
    doc.add_paragraph()
    info = doc.add_paragraph(f"元書籍: 『{book.source_title}』  著: {book.source_author}")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # ─── まえがき ───
    doc.add_heading("まえがき", level=1)
    _add_markdown_paragraphs(doc, book.foreword)
    doc.add_page_break()

    # ─── 各章 ───
    for ch in book.chapters:
        doc.add_heading(f"第{ch.number}章　{ch.title}", level=1)
        _add_markdown_paragraphs(doc, ch.content)
        doc.add_page_break()

    # ─── あとがき ───
    doc.add_heading("あとがき", level=1)
    _add_markdown_paragraphs(doc, book.afterword)

    doc.save(str(out_path))
    console.print(f"[bold green]  ✅ .docx 生成完了: {out_path}[/bold green]")
    return out_path
